import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from agentlang.context.tool_context import ToolContext
from agentlang.tools.tool_result import ToolResult
from app.tools.slide_template_builder import CreateSlideTemplateProject, PrepareSlideTemplateSource, SelectSlideTemplateSourceAssets
from app.tools.slide_template_builder.create_project import CreateSlideTemplateProjectParams
from app.tools.slide_template_builder.models import (
    TemplateDescriptionText,
    TemplateLabelText,
    TemplateSlideParams,
    TemplateSourceParams,
)
from app.tools.slide_template_builder.source_analysis import PrepareSlideTemplateSourceParams
from app.tools.slide_template_builder.source_asset_selection import (
    SelectSlideTemplateSourceAssetsParams,
    SourceAssetDecisionParams,
)
from app.utils.async_file_utils import async_read_json


def _workspace(tmp_path: Path) -> Path:
    workspace = tmp_path / ".workspace"
    workspace.mkdir(exist_ok=True)
    return workspace


def _params() -> CreateSlideTemplateProjectParams:
    return CreateSlideTemplateProjectParams(
        template_id="SLIDE-navy-quarterly-review",
        category_code="PPT-CATE-business-report",
        tag_codes=[
            "purpose-quarterly-review",
            "style-business",
            "color-navy",
            "layout-dashboard",
            "mood-professional",
            "language-english",
        ],
        label=TemplateLabelText(zh_CN="深蓝经营复盘", en_US="Navy Business Review"),
        description=TemplateDescriptionText(
            zh_CN="适合经营复盘。", en_US="A reusable operating review template."
        ),
        source=TemplateSourceParams(kind="derived", method="visual-reconstruction", file="source.pptx"),
        slides=[
            TemplateSlideParams(
                file="slides/cover.html",
                title="Business Review",
                layout="cover",
                description="Opening cover.",
            )
        ],
    )


def _write_source_pptx(path: Path, image_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(0.8), Inches(6), Inches(1))
    textbox.text = "Quarterly review"
    textbox.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    textbox.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    shape = slide.shapes.add_shape(1, Inches(1), Inches(2), Inches(3), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x10, 0x2A, 0x43)
    slide.shapes.add_picture(str(image_path), Inches(7), Inches(1), width=Inches(4), height=Inches(3))
    presentation.save(path)
    package_only = BytesIO()
    Image.new("RGB", (64, 64), "#4f86c6").save(package_only, format="PNG")
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("ppt/media/package-only.png", package_only.getvalue())


async def _prepare_with_assets(tmp_path: Path) -> tuple[Path, Path, dict]:
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    created = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert created.ok, created.content
    template_dir = Path(created.data["template_directory"])
    page_image = tmp_path / "source-page.png"
    embedded_image = tmp_path / "embedded.png"
    source_pptx = tmp_path / "source.pptx"
    Image.new("RGB", (1920, 1080), "#102a43").save(page_image)
    Image.new("RGB", (120, 80), "#e6b566").save(embedded_image)
    _write_source_pptx(source_pptx, embedded_image)

    prepare_tool = PrepareSlideTemplateSource(base_dir=_workspace(tmp_path))
    prepared = await prepare_tool.execute(
        ToolContext(tool_name=prepare_tool.name, arguments={}),
        PrepareSlideTemplateSourceParams(
            template_dir=created.data["template_id"],
            page_images=[str(page_image)],
            source_presentation=str(source_pptx),
        ),
    )
    assert prepared.ok, prepared.content
    work_dir = Path(created.data["working_directory"])
    manifest_path = work_dir / "source-analysis.json"
    return template_dir, work_dir, await async_read_json(manifest_path)


@pytest.mark.asyncio
async def test_select_source_assets_copies_only_reviewed_reusable_candidates(tmp_path):
    template_dir, work_dir, manifest = await _prepare_with_assets(tmp_path)
    catalog = manifest["asset_catalog"]
    decisions = [
        SourceAssetDecisionParams(
            asset=catalog[0]["asset"],
            decision="reuse",
            asset_type="illustration",
            quality="high",
            rationale="The source artwork defines the visual anchor and will be used on the cover.",
            target_name="cover-illustration.png",
            intended_role="Cover visual anchor.",
        ),
        SourceAssetDecisionParams(
            asset=catalog[1]["asset"],
            decision="discard",
            asset_type="other",
            quality="poor",
            rationale="The package-only image is not used by any source slide and adds no reusable value.",
        ),
    ]
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    result = await select_tool.execute(
        ToolContext(tool_name=select_tool.name, arguments={}),
        SelectSlideTemplateSourceAssetsParams(template_dir=template_dir.name, decisions=decisions),
    )

    assert result.ok, result.content
    final_assets = list((template_dir / "images/source-assets").glob("*"))
    assert [path.name for path in final_assets] == ["cover-illustration.png"]
    updated = await async_read_json(work_dir / "source-analysis.json")
    selected = next(item for item in updated["asset_catalog"] if item["decision"] == "reuse")
    discarded = next(item for item in updated["asset_catalog"] if item["decision"] == "discard")
    assert selected["final_file"] == "images/source-assets/cover-illustration.png"
    assert discarded["final_file"] is None
    assert result.data["selected_count"] == 1
    assert result.data["discarded_count"] == 1


@pytest.mark.asyncio
async def test_select_source_assets_accepts_public_asset_ids(tmp_path):
    template_dir, work_dir, manifest = await _prepare_with_assets(tmp_path)
    assert manifest["asset_catalog"][0]["id"] == "asset-01"
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    result = await select_tool.execute(
        ToolContext(tool_name=select_tool.name, arguments={}),
        SelectSlideTemplateSourceAssetsParams(
            template_dir=template_dir.name,
            decisions=[
                SourceAssetDecisionParams(
                    asset="asset-01",
                    decision="reuse",
                    asset_type="illustration",
                    quality="high",
                    rationale="Keep the public-id artwork as the cover visual anchor.",
                    target_name="cover-illustration.png",
                    intended_role="Cover visual anchor.",
                ),
                SourceAssetDecisionParams(
                    asset="asset-02",
                    decision="discard",
                    asset_type="other",
                    quality="poor",
                    rationale="The unused package image should not enter the template.",
                ),
            ],
        ),
    )

    assert result.ok, result.content
    updated = await async_read_json(work_dir / "source-analysis.json")
    selected = next(item for item in updated["asset_catalog"] if item["id"] == "asset-01")
    assert selected["final_file"] == "images/source-assets/cover-illustration.png"


@pytest.mark.asyncio
async def test_select_source_assets_accepts_partial_decisions_then_completes(tmp_path):
    template_dir, work_dir, manifest = await _prepare_with_assets(tmp_path)
    first_asset, second_asset = [item["asset"] for item in manifest["asset_catalog"]]
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    first = await select_tool.execute(
        ToolContext(tool_name=select_tool.name, arguments={}),
        SelectSlideTemplateSourceAssetsParams(
            template_dir=template_dir.name,
            decisions=[
                SourceAssetDecisionParams(
                    asset=first_asset,
                    decision="reuse",
                    asset_type="illustration",
                    quality="high",
                    rationale="Keep the source artwork as the cover visual anchor.",
                    target_name="cover-illustration.png",
                    intended_role="Cover visual anchor.",
                )
            ],
        ),
    )

    assert first.ok, first.content
    assert first.data["pending_count"] == 1
    assert "asset-02" in first.content
    updated = await async_read_json(work_dir / "source-analysis.json")
    assert next(item["decision"] for item in updated["asset_catalog"] if item["asset"] == first_asset) == "reuse"
    assert next(item["decision"] for item in updated["asset_catalog"] if item["asset"] == second_asset) == "pending"

    second = await select_tool.execute(
        ToolContext(tool_name=select_tool.name, arguments={}),
        SelectSlideTemplateSourceAssetsParams(
            template_dir=template_dir.name,
            decisions=[
                SourceAssetDecisionParams(
                    asset=second_asset,
                    decision="discard",
                    asset_type="other",
                    quality="poor",
                    rationale="The package-only image is unused and should not enter the template.",
                )
            ],
        ),
    )

    assert second.ok, second.content
    assert second.data["pending_count"] == 0
    assert second.data["selected_count"] == 1
    assert second.data["discarded_count"] == 1
    assert [path.name for path in (template_dir / "images/source-assets").glob("*")] == ["cover-illustration.png"]


@pytest.mark.asyncio
async def test_select_source_assets_lists_catalog_ids_for_unknown_assets(tmp_path):
    template_dir, _work_dir, manifest = await _prepare_with_assets(tmp_path)
    catalog_ids = [item["id"] for item in manifest["asset_catalog"]]
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    result = await select_tool.execute(
        ToolContext(tool_name=select_tool.name, arguments={}),
        SelectSlideTemplateSourceAssetsParams(
            template_dir=template_dir.name,
            decisions=[
                SourceAssetDecisionParams(
                    asset="source-assets/麦吉AI助理培训_image_0001.webp",
                    decision="reuse",
                    asset_type="illustration",
                    quality="high",
                    rationale="Keep the mascot as a reusable visual element.",
                    target_name="mascot-standing.webp",
                    intended_role="Brand mascot.",
                )
            ],
        ),
    )

    assert not result.ok
    assert "unknown assets" in result.content
    assert "expected catalog asset ids" in result.content
    assert catalog_ids[0] in result.content
    assert not (template_dir / "images/source-assets").exists()


def test_source_asset_decision_rejects_poor_reuse():
    with pytest.raises(ValueError, match="poor-quality assets cannot be reused"):
        SourceAssetDecisionParams(
            asset="source-assets/asset.png",
            decision="reuse",
            asset_type="illustration",
            quality="poor",
            rationale="This is too blurry to use.",
            target_name="asset.png",
            intended_role="Cover artwork.",
        )


def test_source_asset_decision_accepts_catalog_id_alias():
    decision = SourceAssetDecisionParams.model_validate({
        "id": "asset-01",
        "decision": "discard",
        "asset_type": "photo",
        "quality": "poor",
        "rationale": "Unused low-resolution crop.",
    })
    assert decision.asset == "asset-01"
    assert decision.decision == "discard"


def test_select_source_assets_params_reject_catalog_id_strings():
    with pytest.raises(ValueError, match="complete object"):
        SelectSlideTemplateSourceAssetsParams.model_validate({
            "template_dir": "navy-review",
            "decisions": ["asset-01", "asset-02"],
        })


@pytest.mark.asyncio
async def test_select_source_assets_explains_incomplete_decision_objects(tmp_path):
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    result = await select_tool(
        ToolContext(tool_name=select_tool.name, arguments={}),
        template_dir="navy-review",
        decisions=[{"id": f"asset-{index:02d}"} for index in range(1, 13)],
    )

    assert not result.ok
    assert "complete object" in result.content
    assert '"asset": "asset-01"' in result.content
    assert "decisions.0.asset" not in result.content
    assert "decisions.0.decision" not in result.content


@pytest.mark.asyncio
async def test_select_source_assets_accepts_public_id_alias_through_tool_call(tmp_path):
    template_dir, work_dir, _manifest = await _prepare_with_assets(tmp_path)
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    result = await select_tool(
        ToolContext(tool_name=select_tool.name, arguments={}),
        template_dir=template_dir.name,
        decisions=[
            {
                "id": "asset-01",
                "decision": "reuse",
                "asset_type": "illustration",
                "quality": "high",
                "rationale": "Keep the public-id artwork as the cover visual anchor.",
                "target_name": "cover-illustration.png",
                "intended_role": "Cover visual anchor.",
            },
            {
                "id": "asset-02",
                "decision": "discard",
                "asset_type": "other",
                "quality": "poor",
                "rationale": "The unused package image should not enter the template.",
            },
        ],
    )

    assert result.ok, result.content
    updated = await async_read_json(work_dir / "source-analysis.json")
    selected = next(item for item in updated["asset_catalog"] if item["id"] == "asset-01")
    assert selected["final_file"] == "images/source-assets/cover-illustration.png"


def test_source_asset_decision_normalizes_natural_aliases():
    discarded = SourceAssetDecisionParams.model_validate({
        "id": "asset-15",
        "action": "discard",
        "type": "logo",
        "quality": "low",
        "note": "iSlide promotional watermark.",
    })
    assert discarded.asset == "asset-15"
    assert discarded.decision == "discard"
    assert discarded.quality == "poor"
    assert discarded.rationale == "iSlide promotional watermark."

    reused = SourceAssetDecisionParams.model_validate({
        "asset": "asset-01",
        "decision": "reuse",
        "asset_type": "photograph",
        "quality": "high",
        "rationale": "Eiffel Tower background photo for cover page.",
        "target_name": "eiffel-tower-sky.jpg",
    })
    assert reused.asset_type == "photo"
    assert reused.intended_role == "Eiffel Tower background photo for cover page."


def test_source_asset_decision_defaults_reuse_role_and_keeps_missing_target_name():
    decision = SourceAssetDecisionParams.model_validate({
        "asset": "asset-02",
        "decision": "reuse",
        "asset_type": "illustration",
        "quality": "high",
        "rationale": "Gold trophy for the agenda page.",
    })
    assert decision.intended_role == "Gold trophy for the agenda page."
    assert decision.target_name is None


@pytest.mark.asyncio
async def test_select_source_assets_rewrites_wrong_target_extension(tmp_path):
    template_dir, work_dir, _manifest = await _prepare_with_assets(tmp_path)
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    result = await select_tool(
        ToolContext(tool_name=select_tool.name, arguments={}),
        template_dir=template_dir.name,
        decisions=[
            {
                "asset": "asset-01",
                "decision": "reuse",
                "asset_type": "photo",
                "quality": "high",
                "rationale": "Keep the cover photo.",
                "target_name": "eiffel-tower-sky.jpg",
                "intended_role": "Cover background.",
            },
            {
                "asset": "asset-02",
                "decision": "discard",
                "asset_type": "other",
                "quality": "poor",
                "rationale": "Unused package image.",
            },
        ],
    )

    assert result.ok, result.content
    updated = await async_read_json(work_dir / "source-analysis.json")
    selected = next(item for item in updated["asset_catalog"] if item["id"] == "asset-01")
    assert selected["target_name"] == "eiffel-tower-sky.png"
    assert selected["final_file"] == "images/source-assets/eiffel-tower-sky.png"
    assert (template_dir / "images/source-assets/eiffel-tower-sky.png").exists()


@pytest.mark.asyncio
async def test_select_source_assets_fills_missing_reuse_filename(tmp_path):
    template_dir, work_dir, _manifest = await _prepare_with_assets(tmp_path)
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    result = await select_tool(
        ToolContext(tool_name=select_tool.name, arguments={}),
        template_dir=template_dir.name,
        decisions=[
            {
                "asset": "asset-01",
                "decision": "reuse",
                "asset_type": "illustration",
                "quality": "high",
                "rationale": "Keep the source artwork as the cover visual anchor.",
            },
            {
                "asset": "asset-02",
                "decision": "discard",
                "asset_type": "other",
                "quality": "low",
                "rationale": "Unused package image.",
            },
        ],
    )

    assert result.ok, result.content
    updated = await async_read_json(work_dir / "source-analysis.json")
    selected = next(item for item in updated["asset_catalog"] if item["id"] == "asset-01")
    assert selected["target_name"] == "asset-01.png"
    assert selected["final_file"] == "images/source-assets/asset-01.png"


@pytest.mark.asyncio
async def test_select_source_assets_remark_shows_template_name_not_path():
    select_tool = SelectSlideTemplateSourceAssets(base_dir="/tmp")
    arguments = {"template_dir": "/mnt/agfs-vol/agfs/magicfs/project/SLIDE-奥运体育赛事"}
    before = await select_tool.get_before_tool_call_friendly_action_and_remark(
        select_tool.name,
        ToolContext(tool_name=select_tool.name, arguments=arguments),
        arguments,
    )
    after = await select_tool.get_after_tool_call_friendly_action_and_remark(
        select_tool.name,
        ToolContext(tool_name=select_tool.name, arguments=arguments),
        ToolResult(
            content="ok",
            data={"template_directory": "/mnt/agfs-vol/agfs/magicfs/project/SLIDE-奥运体育赛事"},
        ),
        0.1,
        arguments,
    )

    assert "SLIDE-奥运体育赛事" in before["remark"]
    assert "SLIDE-奥运体育赛事" in after["remark"]
    assert "/mnt/" not in before["remark"]
    assert "/mnt/" not in after["remark"]
    assert "magicfs" not in before["remark"]
    assert "magicfs" not in after["remark"]
