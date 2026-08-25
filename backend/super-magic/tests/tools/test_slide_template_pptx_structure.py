import zipfile
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.tools.slide_template_builder.pptx_structure import extract_pptx_structure


def _write_picture_pptx(path: Path, image_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), width=Inches(3))
    presentation.save(path)


def _rewrite_png_content_type(pptx_path: Path, content_type: str) -> None:
    source_bytes = BytesIO(pptx_path.read_bytes())
    output = BytesIO()
    with zipfile.ZipFile(source_bytes) as source, zipfile.ZipFile(output, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(b'ContentType="image/png"', f'ContentType="{content_type}"'.encode("utf-8"))
            target.writestr(info, data)
    pptx_path.write_bytes(output.getvalue())


def test_extract_pptx_structure_keeps_unregistered_picture_parts(tmp_path):
    image_path = tmp_path / "mascot.png"
    pptx_path = tmp_path / "source.pptx"
    Image.new("RGB", (80, 80), "#ff8800").save(image_path)
    _write_picture_pptx(pptx_path, image_path)
    _rewrite_png_content_type(pptx_path, "image/svg+xml")

    structure = extract_pptx_structure(pptx_path)

    assert structure.metadata["slide_count"] == 1
    assert structure.metadata["source_asset_count"] >= 1
    assert any(asset.relative_path.startswith("source-assets/") for asset in structure.assets)
    picture_shapes = [
        shape
        for shape in structure.metadata["slides"][0]["shapes"]
        if shape.get("type") == "PICTURE"
    ]
    assert picture_shapes
    assert picture_shapes[0].get("image", {}).get("asset")
