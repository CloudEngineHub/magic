import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from agentlang.context.tool_context import ToolContext
from app.tools.slide_template_builder import CreateSlideTemplateProject, PrepareSlideTemplateSource
from app.tools.slide_template_builder.create_project import CreateSlideTemplateProjectParams
from app.tools.slide_template_builder.models import (
    TemplateCanvas,
    TemplateDescriptionText,
    TemplateLabelText,
    TemplateSlideParams,
    TemplateSourceParams,
)
from app.tools.slide_template_builder.pptx_structure import extract_pptx_structure
from app.tools.slide_template_builder.preview import (
    _SLIDE_INSPECTION_SCRIPT,
    _create_collage,
    render_slide_previews,
)
from app.tools.slide_template_builder.source_analysis import PrepareSlideTemplateSourceParams
from magic_use.config import BrowserRuntimeConfig


def _params() -> CreateSlideTemplateProjectParams:
    return CreateSlideTemplateProjectParams(
        template_id="SLIDE-square-source",
        category_code="PPT-CATE-business-report",
        tag_codes=["style-business"],
        label=TemplateLabelText(zh_CN="方形模板", en_US="Square Template"),
        description=TemplateDescriptionText(zh_CN="保留方形画布。", en_US="Preserves a square canvas."),
        source=TemplateSourceParams(kind="derived", method="visual-reconstruction", file="source.pptx"),
        slides=[
            TemplateSlideParams(
                file="slides/cover.html",
                title="Square Cover",
                layout="cover",
                description="Square opening cover.",
            )
        ],
    )


def _write_square_pptx(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(10)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


class FakeBrowserClient:
    def __init__(self, inspection: dict, screenshot: bytes):
        self.inspection = inspection
        self.screenshot_bytes = screenshot
        self.navigated_urls = []
        self.closed = False

    async def list_pages(self):
        return [SimpleNamespace(id="page-1")]

    async def navigate(self, page_id, url):
        assert page_id == "page-1"
        self.navigated_urls.append(url)

    async def evaluate(self, page_id, expression, argument=None):
        assert page_id == "page-1"
        if argument is not None:
            assert "URL.createObjectURL" in expression
            return "blob:https://browser.example/rendered-page"
        if "document.fonts" in expression:
            return None
        return self.inspection

    async def screenshot(self, page_id):
        assert page_id == "page-1"
        return SimpleNamespace(image=self.screenshot_bytes)

    async def close(self):
        self.closed = True


@pytest.mark.asyncio
async def test_render_slide_previews_uses_fixed_canvas_browser_sdk_and_blob_url(tmp_path):
    template_dir = tmp_path / "SLIDE-square-source"
    slide_path = template_dir / "slides/封面.html"
    from app.utils.async_file_utils import async_mkdir, async_write_text

    await async_mkdir(slide_path.parent, parents=True)
    await async_write_text(slide_path, '<!doctype html><main class="slide"></main>')
    inspection = {
        "scrollWidth": 1080,
        "scrollHeight": 1080,
        "bodyScrollWidth": 1080,
        "bodyScrollHeight": 1080,
        "brokenImages": [],
        "slideBounds": {"width": 1080, "height": 1080},
        "viewport": {"width": 1080, "height": 1080},
    }
    screenshot = BytesIO()
    Image.new("RGB", (1080, 1080), "#102a43").save(screenshot, format="PNG")
    browser = FakeBrowserClient(inspection, screenshot.getvalue())
    base_config = BrowserRuntimeConfig()

    with (
        patch(
            "app.tools.slide_template_builder.preview.BrowserConfigAdapter.build_playwright",
            new=AsyncMock(return_value=base_config),
        ) as build_playwright_mock,
        patch(
            "app.tools.slide_template_builder.preview.create_browser",
            new=AsyncMock(return_value=browser),
        ) as create_browser_mock,
    ):
        captures = await render_slide_previews(
            workspace_dir=tmp_path,
            template_dir=template_dir,
            slides=[{"file": "slides/封面.html"}],
            canvas=TemplateCanvas(width=1080, height=1080),
        )

    build_playwright_mock.assert_awaited_once_with(
        str(tmp_path),
        viewport_width=1080,
        viewport_height=1080,
        device_scale_factor=1,
    )
    assert create_browser_mock.await_args.args[0] is base_config
    assert len(captures) == 1
    assert captures[0].inspection == inspection
    assert browser.navigated_urls == ["blob:https://browser.example/rendered-page"]
    assert browser.closed is True
    assert "document.querySelector('.slide-container, .slide')" in _SLIDE_INSPECTION_SCRIPT


@pytest.mark.asyncio
async def test_create_project_preserves_square_source_canvas(tmp_path):
    params = _params()
    params.source.canvas.width = 1080
    params.source.canvas.height = 1080

    tool = CreateSlideTemplateProject(base_dir=tmp_path)
    result = await tool.execute(ToolContext(tool_name=tool.name, arguments={}), params)

    assert result.ok, result.content
    template_dir = Path(result.data["template_directory"])
    metadata = json.loads((template_dir / "template.json").read_text())
    assert metadata["source"]["canvas"] == {"width": 1080, "height": 1080}
    assert result.data["canvas_contract"]["width"] == 1080
    assert result.data["canvas_contract"]["height"] == 1080
    scaffold = (template_dir / "slides/cover.html").read_text()
    assert 'content="width=1080, height=1080"' in scaffold
    assert '<main class="slide-container"' in scaffold
    assert 'data-width="1080" data-height="1080"' in scaffold


@pytest.mark.asyncio
async def test_prepare_source_updates_canvas_when_rendered_pages_are_square(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=tmp_path)
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok

    page_image = tmp_path / "square-source-page.png"
    Image.new("RGB", (1080, 1080), "#102a43").save(page_image)
    source_pptx = tmp_path / "square-source.pptx"
    _write_square_pptx(source_pptx)
    prepare_tool = PrepareSlideTemplateSource(base_dir=tmp_path)
    result = await prepare_tool.execute(
        ToolContext(tool_name=prepare_tool.name, arguments={}),
        PrepareSlideTemplateSourceParams(
            template_dir=create_result.data["template_id"],
            page_images=[str(page_image)],
            source_presentation=str(source_pptx),
        ),
    )

    assert result.ok, result.content
    assert result.data["canvas"] == {"width": 1920, "height": 1920}
    metadata = json.loads((Path(create_result.data["template_directory"]) / "template.json").read_text())
    assert metadata["source"]["canvas"] == {"width": 1920, "height": 1920}
    assert "Canvas: 1920x1920" in result.content
    assert "preserve the source presentation aspect ratio" in result.content


@pytest.mark.asyncio
async def test_prepare_source_rejects_rendered_ratio_that_differs_from_pptx(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=tmp_path)
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok

    page_image = tmp_path / "widescreen-page.png"
    Image.new("RGB", (1920, 1080), "#102a43").save(page_image)
    source_pptx = tmp_path / "square-source.pptx"
    _write_square_pptx(source_pptx)

    prepare_tool = PrepareSlideTemplateSource(base_dir=tmp_path)
    result = await prepare_tool.execute(
        ToolContext(tool_name=prepare_tool.name, arguments={}),
        PrepareSlideTemplateSourceParams(
            template_dir=create_result.data["template_id"],
            page_images=[str(page_image)],
            source_presentation=str(source_pptx),
        ),
    )

    assert not result.ok
    assert "does not match the PPTX canvas 1920x1920" in result.content


def test_pptx_structure_preserves_square_slide_ratio(tmp_path):
    source_pptx = tmp_path / "square-source.pptx"
    _write_square_pptx(source_pptx)

    structure = extract_pptx_structure(source_pptx)

    assert structure.metadata["normalized_canvas"] == {"width": 1920, "height": 1920}


def test_square_collage_does_not_stretch_to_widescreen():
    square = BytesIO()
    Image.new("RGB", (1080, 1080), "#123456").save(square, format="PNG")

    collage = _create_collage([square.getvalue()], TemplateCanvas(width=1080, height=1080))

    with Image.open(BytesIO(collage)) as image:
        assert image.size == (424, 424)
