import json
import os
import re
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from agentlang.context.tool_context import ToolContext
from app.tools.slide_template_builder import (
    CreateSlideTemplateProject,
    FinalizeSlideTemplate,
    PrepareSlideTemplateSource,
    SelectSlideTemplateSourceAssets,
)
from app.tools.slide_template_builder.common import (
    WORK_DIR_NAME,
    background_variants,
    cleanup_template_intermediates,
    describe_background_variants,
    evict_stale_template_working_directories,
    generate_semantic_directory_name,
    recommended_review_indexes,
    resolve_template_working_directory,
    unreviewed_background_variants,
    validate_template_id,
)
from app.tools.slide_template_builder.create_project import CreateSlideTemplateProjectParams
from app.tools.slide_template_builder.finalize_project import FinalizeSlideTemplateParams
from app.tools.slide_template_builder.models import (
    TemplateDescriptionText,
    TemplateLabelText,
    TemplateSlideParams,
    TemplateSourceParams,
)
from app.tools.slide_template_builder.preview import _create_collage, _resize_png
from app.tools.slide_template_builder.source_analysis import PrepareSlideTemplateSourceParams
from app.tools.slide_template_builder.source_asset_selection import (
    SelectSlideTemplateSourceAssetsParams,
    SourceAssetDecisionParams,
)
from app.tools.slide_template_builder.source_images import probe_page_background
from app.tools.slide_template_builder.source_review import (
    parse_source_review_payload,
    render_page_digest,
)
from app.utils.async_file_utils import async_read_json, async_scandir, async_write_json, async_write_text


def _workspace(tmp_path: Path) -> Path:
    workspace = tmp_path / ".workspace"
    workspace.mkdir(exist_ok=True)
    return workspace


def _params():
    return CreateSlideTemplateProjectParams(
        template_id="SLIDE-navy-quarterly-review",
        category_code="PPT-CATE-business-report",
        tag_codes=[
            "purpose-quarterly-review",
            "style-business",
            "color-navy",
            "layout-dashboard",
            "mood-professional",
            "language-english",
        ],
        label=TemplateLabelText(zh_CN="深蓝经营复盘", en_US="Navy Business Review"),
        description=TemplateDescriptionText(zh_CN="适合经营复盘。", en_US="A reusable operating review template."),
        source=TemplateSourceParams(kind="derived", method="visual-reconstruction", file="source.pptx"),
        slides=[
            TemplateSlideParams(
                file="slides/cover.html",
                title="Business Review",
                layout="cover",
                description="Opening cover.",
            )
        ],
    )


def _write_source_pptx(path: Path, image_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(0.8), Inches(6), Inches(1))
    textbox.text = "Quarterly review"
    textbox.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    textbox.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    shape = slide.shapes.add_shape(1, Inches(1), Inches(2), Inches(3), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x10, 0x2A, 0x43)
    slide.shapes.add_picture(str(image_path), Inches(7), Inches(1), width=Inches(4), height=Inches(3))
    presentation.save(path)
    package_only = BytesIO()
    Image.new("RGB", (64, 64), "#4f86c6").save(package_only, format="PNG")
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("ppt/media/package-only.png", package_only.getvalue())


@pytest.mark.asyncio
async def test_create_project_puts_standard_files_at_workspace_root(tmp_path):
    tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    result = await tool.execute(ToolContext(tool_name=tool.name, arguments={}), _params())

    assert result.ok
    template_dir = Path(result.data["template_directory"])
    assert result.data["template_directory"] == str(template_dir)
    assert result.data["requested_template_id"] == "SLIDE-navy-quarterly-review"
    assert re.fullmatch(r"SLIDE-navy-quarterly-review-[0-9A-F]{6}", result.data["template_id"])
    assert template_dir.name == "SLIDE-navy-business-review"
    assert (template_dir / "template.json").exists()
    assert (template_dir / "visual-spec.md").exists()
    assert (template_dir / "theme.css").exists()
    assert (template_dir / "images").is_dir()
    assert (template_dir / "slides/cover.html").exists()
    assert not (template_dir / WORK_DIR_NAME).exists()
    work_dir = Path(result.data["working_directory"])
    assert work_dir.is_dir()
    assert work_dir == resolve_template_working_directory(template_dir, result.data["template_id"])
    with pytest.raises(ValueError):
        work_dir.resolve().relative_to(template_dir.parent.resolve())
    work_dir.resolve().relative_to((tmp_path / ".runtime").resolve())
    converter_output_dir = Path(result.data["converter_output_dir"])
    assert converter_output_dir.is_dir()
    assert converter_output_dir.parent == work_dir
    converter_output_dir.resolve().relative_to((tmp_path / ".runtime").resolve())
    with pytest.raises(ValueError):
        converter_output_dir.resolve().relative_to(template_dir.parent.resolve())
    assert str(converter_output_dir) in result.content
    assert "Do not write PDF, page images, or converter dumps into the workspace" in result.content
    assert result.data["canvas_contract"] == {
        "selector": ".slide-container",
        "width_attribute": "data-width",
        "height_attribute": "data-height",
        "width": 1920,
        "height": 1080,
    }
    assert "main.slide-container" in result.content
    assert "template.json.source.canvas" in result.content
    assert "Do not replace `template.json`" in result.content
    assert "the template is not complete until finalization succeeds" in result.content
    assert "slide-template-scaffold" in (template_dir / "theme.css").read_text()
    scaffold = (template_dir / "slides/cover.html").read_text()
    assert '<main class="slide-container"' in scaffold
    assert 'data-width="1920" data-height="1080"' in scaffold
    metadata = json.loads((template_dir / "template.json").read_text())
    assert metadata["template_id"] == result.data["template_id"]
    assert "directory_name" not in metadata
    assert metadata["category_code"] == "PPT-CATE-business-report"
    assert metadata["tag_codes"] == [
        "purpose-quarterly-review",
        "style-business",
        "color-navy",
        "layout-dashboard",
        "mood-professional",
        "language-english",
    ]


def test_generate_semantic_directory_name_uses_english_label_and_sanitizes_punctuation():
    name = generate_semantic_directory_name(
        {"zh_CN": "增长 / 路线图：2026", "en_US": "Growth Roadmap: 2026"},
        ["language-english"],
    )

    assert name == "SLIDE-growth-roadmap-2026"


def test_generate_semantic_directory_name_uses_chinese_label():
    name = generate_semantic_directory_name(
        {"zh_CN": "增长路线图", "en_US": "Growth Roadmap"},
        ["language-chinese"],
    )

    assert name == "SLIDE-增长路线图"


@pytest.mark.asyncio
async def test_create_project_rejects_an_existing_semantic_directory(tmp_path):
    tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    first = await tool.execute(ToolContext(tool_name=tool.name, arguments={}), _params())
    second = await tool.execute(ToolContext(tool_name=tool.name, arguments={}), _params())

    assert first.ok, first.content
    assert not second.ok
    assert "Semantic template directory already exists" in second.content


@pytest.mark.asyncio
async def test_create_project_does_not_leave_a_directory_when_slide_plan_is_invalid(tmp_path):
    params = _params()
    params.slides = [params.slides[0], params.slides[0]]
    tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))

    result = await tool.execute(ToolContext(tool_name=tool.name, arguments={}), params)

    assert not result.ok
    assert "Slide file paths must be unique" in result.content
    created_entries = await async_scandir(_workspace(tmp_path))
    assert not [entry for entry in created_entries if entry.name == "SLIDE-navy-business-review"]


@pytest.mark.asyncio
async def test_create_project_rejects_unknown_taxonomy_codes(tmp_path):
    params = _params()
    params.category_code = "PPT-CATE-not-a-real-category"
    tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    result = await tool.execute(ToolContext(tool_name=tool.name, arguments={}), params)
    assert not result.ok
    assert "active bundled PPT category" in result.content

    params = _params()
    params.tag_codes = ["purpose-quarterly-review", "tag-does-not-exist"]
    result = await tool.execute(ToolContext(tool_name=tool.name, arguments={}), params)
    assert not result.ok
    assert "unknown bundled codes" in result.content


@pytest.mark.asyncio
async def test_prepare_source_creates_multiscale_evidence_and_pptx_facts(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok
    template_dir = Path(create_result.data["template_directory"])

    page_image = tmp_path / "source-page.png"
    Image.new("RGB", (1920, 1080), "#102a43").save(page_image)
    source_pptx = tmp_path / "source.pptx"
    embedded_image = tmp_path / "embedded.png"
    Image.new("RGB", (120, 80), "#e6b566").save(embedded_image)
    _write_source_pptx(source_pptx, embedded_image)

    prepare_tool = PrepareSlideTemplateSource(base_dir=_workspace(tmp_path))
    result = await prepare_tool.execute(
        ToolContext(tool_name=prepare_tool.name, arguments={}),
        PrepareSlideTemplateSourceParams(
            template_dir=create_result.data["template_id"],
            page_images=[str(page_image)],
            source_presentation=str(source_pptx),
        ),
    )

    assert result.ok, result.content
    work_dir = Path(create_result.data["working_directory"])
    analysis_root = work_dir / "source-analysis"
    assert not (template_dir / WORK_DIR_NAME).exists()
    assert result.data["working_directory"] == str(work_dir)
    assert (analysis_root / "full/slide-001.png").exists()
    assert len(list((analysis_root / "crops/slide-001").glob("*.png"))) == 9
    assert (analysis_root / "contact-sheet.png").exists()
    final_source_assets_dir = template_dir / "images/source-assets"
    assert not final_source_assets_dir.exists()
    structure = json.loads((analysis_root / "pptx-structure.json").read_text())
    assert structure["slide_count"] == 1
    assert structure["source_asset_count"] == 2
    assert any(shape.get("image") for shape in structure["slides"][0]["shapes"])
    assert list((analysis_root / "source-assets").glob("*"))
    manifest = json.loads((work_dir / "source-analysis.json").read_text())
    assert manifest["slides"][0]["full_image"] == "source-analysis/full/slide-001.png"
    assert manifest["slides"][0]["visual_analysis"]["status"] == "pending"
    assert len(manifest["asset_catalog"]) == 2
    assert [item["id"] for item in manifest["asset_catalog"]] == ["asset-01", "asset-02"]
    assert "asset-01" in result.content
    assert result.data["recommended_review_indexes"] == [1]
    assert all(item["decision"] == "pending" for item in manifest["asset_catalog"])
    assert all(item["final_file"] is None for item in manifest["asset_catalog"])
    assert result.data["detail_crop_count"] == 9
    assert result.data["extracted_asset_count"] == 2
    assert len(result.data["asset_analysis_batches"]) == 1

def test_background_probe_reads_flat_and_image_backed_pages():
    flat = probe_page_background(Image.new("RGB", (1920, 1080), "#102a43"))
    assert flat.color == "#102a43"
    assert flat.tone == "dark"
    assert flat.uniform

    light = probe_page_background(Image.new("RGB", (1920, 1080), "#f5f7fa"))
    assert light.tone == "light"
    assert light.uniform

    row = Image.new("RGB", (1920, 1))
    for x in range(1920):
        ratio = x / 1919
        row.putpixel((x, 0), (
            round(0x10 + (0xF7 - 0x10) * ratio),
            round(0x2A + (0xF8 - 0x2A) * ratio),
            round(0x43 + (0xFA - 0x43) * ratio),
        ))
    assert not probe_page_background(row.resize((1920, 1080))).uniform


def test_background_variants_group_pages_and_drive_review_sampling():
    dark = {"color": "#102a43", "tone": "dark", "uniform": True}
    light = {"color": "#f7f8fa", "tone": "light", "uniform": True}
    slides = [
        {"source_index": index, "background_probe": dark if index == 2 else light}
        for index in range(1, 13)
    ]
    variants = background_variants(slides)
    assert [variant["pages"] for variant in variants] == [[1, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12], [2]]
    assert describe_background_variants(variants) == (
        "light #f7f8fa (flat) on page 1, 3-12; dark #102a43 (flat) on page 2"
    )

    # 默认抽样漏掉了唯一的深色页，补充变体代表页后它必须进入复核范围
    assert recommended_review_indexes(12) == [1, 4, 8, 12]
    assert recommended_review_indexes(12, variants) == [1, 2, 4, 8, 12]
    assert unreviewed_background_variants(slides) == variants


def test_review_digest_exposes_measured_background_and_geometry():
    payload = parse_source_review_payload(
        json.dumps({
            "summary": "Light agenda page on a 6% left axis.",
            "background": {"tone": "dark", "base_color": "#000000", "text_color": "#1f2933"},
            "palette": [{"role": "accent", "hex": "#e6b566", "usage": "step numbers"}],
            "typography": [{"role": "page-title", "size_pct_of_height": 6.5, "weight": 700}],
            "composition": {
                "margin_pct": {"top": 8, "right": 6, "bottom": 8, "left": 6},
                "columns": 2,
                "title_block": {"x_pct": 6, "y_pct": 10, "width_pct": 55, "align": "left"},
                "density": "medium",
            },
            "must_preserve": ["Gold hairline under the title"],
        }),
        {"color": "#f7f8fa", "tone": "light", "luminance": 0.9, "uniform": True},
    )

    # 测量值覆盖模型猜测，避免浅色页被写成深色
    assert payload["background"]["base_color"] == "#f7f8fa"
    assert payload["background"]["tone"] == "light"

    digest = render_page_digest(4, payload)
    assert "light #f7f8fa" in digest
    assert "Body text #1f2933" in digest
    assert "margins T8% R6% B8% L6%" in digest
    assert "title block x6% y10% w55% left" in digest
    assert "page-title 6.5% of height weight 700" in digest
    assert "accent #e6b566 (step numbers)" in digest


@pytest.mark.asyncio
async def test_visual_reconstruction_requires_completed_source_analysis(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok
    template_dir = Path(create_result.data["template_directory"])
    await async_write_text(
        template_dir / "visual-spec.md",
        "# Visual Specification\n\nNavy report system.\n",
    )
    await async_write_text(
        template_dir / "theme.css",
        "html, body { margin: 0; width: 1920px; height: 1080px; overflow: hidden; }\n.slide { width: 1920px; height: 1080px; }\n",
    )
    await async_write_text(
        template_dir / "slides/cover.html",
        "<!doctype html><html><head><link rel=\"stylesheet\" href=\"../theme.css\"></head><body><main class=\"slide\"><h1 data-slot=\"title\">Business Review</h1><p data-slot=\"subtitle\">Q4</p><span data-slot=\"meta\">Strategy Team</span></main></body></html>",
    )
    from app.tools.slide_template_builder.common import validate_template_project

    with pytest.raises(Exception, match="completed source analysis"):
        await validate_template_project(template_dir)


@pytest.mark.asyncio
async def test_incomplete_visual_analysis_names_status_and_manifest_path(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok
    template_dir = Path(create_result.data["template_directory"])
    page_image = tmp_path / "source-page.png"
    Image.new("RGB", (1920, 1080), "#102a43").save(page_image)
    prepare_tool = PrepareSlideTemplateSource(base_dir=_workspace(tmp_path))
    prepared = await prepare_tool.execute(
        ToolContext(tool_name=prepare_tool.name, arguments={}),
        PrepareSlideTemplateSourceParams(
            template_dir=create_result.data["template_id"],
            page_images=[str(page_image)],
        ),
    )
    assert prepared.ok, prepared.content
    await async_write_text(template_dir / "visual-spec.md", "# Visual Specification\n\nNavy report system.\n")
    await async_write_text(
        template_dir / "theme.css",
        "html, body { margin: 0; width: 1920px; height: 1080px; overflow: hidden; }\n.slide { width: 1920px; height: 1080px; }\n",
    )
    await async_write_text(
        template_dir / "slides/cover.html",
        "<!doctype html><html><head><link rel=\"stylesheet\" href=\"../theme.css\"></head><body><main class=\"slide\"><h1 data-slot=\"title\">Business Review</h1><p data-slot=\"subtitle\">Q4</p><span data-slot=\"meta\">Strategy Team</span></main></body></html>",
    )
    from app.tools.slide_template_builder.common import validate_template_project

    validated = await validate_template_project(template_dir)
    assert validated.source_analysis_summary is not None
    assert validated.source_analysis_summary["unreviewed_source_indexes"] == [1]
    assert any("review_slide_template_source" in warning for warning in validated.source_analysis_summary["warnings"])


@pytest.mark.asyncio
async def test_validate_project_rejects_mismatched_slide_container_dimensions(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok
    template_dir = Path(create_result.data["template_directory"])
    await async_write_text(template_dir / "visual-spec.md", "# Visual Specification\n\nNavy report system.\n")
    await async_write_text(template_dir / "theme.css", "html, body { margin: 0; }\n")
    await async_write_text(
        template_dir / "slides/cover.html",
        '<!doctype html><html><head><link rel="stylesheet" href="../theme.css"></head>'
        '<body><main class="slide-container" data-width="1080" data-height="1080">'
        '<h1 data-slot="title">Business Review</h1><p data-slot="subtitle">Q4</p>'
        '<span data-slot="meta">Strategy Team</span></main></body></html>',
    )

    from app.tools.slide_template_builder.common import validate_template_project

    with pytest.raises(Exception, match="must match source.canvas 1920x1080"):
        await validate_template_project(template_dir)


@pytest.mark.asyncio
async def test_completed_source_analysis_is_accepted_and_summarized(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok
    template_dir = Path(create_result.data["template_directory"])
    page_image = tmp_path / "source-page.png"
    Image.new("RGB", (1920, 1080), "#102a43").save(page_image)

    prepare_tool = PrepareSlideTemplateSource(base_dir=_workspace(tmp_path))
    prepare_result = await prepare_tool.execute(
        ToolContext(tool_name=prepare_tool.name, arguments={}),
        PrepareSlideTemplateSourceParams(
            template_dir=create_result.data["template_id"],
            page_images=[str(page_image)],
        ),
    )
    assert prepare_result.ok, prepare_result.content

    await async_write_text(template_dir / "visual-spec.md", "# Visual Specification\n\nDark report system.\n")
    await async_write_text(
        template_dir / "theme.css",
        "html, body { margin: 0; width: 1920px; height: 1080px; overflow: hidden; }\n.slide { width: 1920px; height: 1080px; }\n",
    )
    await async_write_text(
        template_dir / "slides/cover.html",
        "<!doctype html><html><head><link rel=\"stylesheet\" href=\"../theme.css\"></head><body><main class=\"slide\"><h1 data-slot=\"title\">Business Review</h1><p data-slot=\"subtitle\">Q4</p><span data-slot=\"meta\">Strategy Team</span></main></body></html>",
    )

    work_dir = Path(create_result.data["working_directory"])
    manifest_path = work_dir / "source-analysis.json"
    manifest = await async_read_json(manifest_path)
    visual = manifest["slides"][0]["visual_analysis"]
    visual.update({
        "status": "complete",
        "summary": "A dark report cover with a strong left-aligned title anchor.",
        "visual_elements": [{"role": "title-anchor", "description": "Large left-aligned title."}],
        "style_tokens": {"accent": "#e6b566", "corner_radius_px": 0},
        "must_preserve": ["Dark navy canvas", "Left title anchor"],
    })
    await async_write_json(manifest_path, manifest, ensure_ascii=False, indent=2)

    from app.tools.slide_template_builder.common import validate_template_project

    validated = await validate_template_project(template_dir)
    assert validated.source_analysis_summary == {
        "completed": True,
        "reconstruction_mode": "hybrid",
        "source_slide_count": 1,
        "reviewed_slide_count": 1,
        "unreviewed_source_indexes": [],
        "background_variants": [
            {
                "key": "dark:flat:1-2-3",
                "color": "#102a43",
                "tone": "dark",
                "uniform": True,
                "pages": [1],
            }
        ],
        "background_map": "dark #102a43 (flat) on page 1",
        "unreviewed_background_variants": [],
        "must_preserve_count": 2,
        "visual_element_count": 1,
        "pptx_structure_extracted": False,
        "source_asset_count": 0,
        "selected_source_asset_count": 0,
        "discarded_source_asset_count": 0,
        "selected_source_files": [],
        "referenced_source_asset_count": 0,
        "warnings": [],
    }


@pytest.mark.asyncio
async def test_hybrid_reconstruction_must_reference_selected_source_asset(tmp_path):
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok
    template_dir = Path(create_result.data["template_directory"])
    page_image = tmp_path / "source-page.png"
    embedded_image = tmp_path / "embedded.png"
    source_pptx = tmp_path / "source.pptx"
    Image.new("RGB", (1920, 1080), "#102a43").save(page_image)
    Image.new("RGB", (120, 80), "#e6b566").save(embedded_image)
    _write_source_pptx(source_pptx, embedded_image)

    prepare_tool = PrepareSlideTemplateSource(base_dir=_workspace(tmp_path))
    prepare_result = await prepare_tool.execute(
        ToolContext(tool_name=prepare_tool.name, arguments={}),
        PrepareSlideTemplateSourceParams(
            template_dir=create_result.data["template_id"],
            page_images=[str(page_image)],
            source_presentation=str(source_pptx),
        ),
    )
    assert prepare_result.ok, prepare_result.content
    work_dir = Path(create_result.data["working_directory"])
    manifest_path = work_dir / "source-analysis.json"
    manifest = await async_read_json(manifest_path)
    manifest["slides"][0]["visual_analysis"].update({
        "status": "complete",
        "summary": "Original image-led report cover.",
        "visual_elements": [{"role": "source-illustration"}],
        "style_tokens": {"canvas": "#102a43"},
        "must_preserve": ["Original illustration"],
    })
    await async_write_json(manifest_path, manifest, ensure_ascii=False, indent=2)
    catalog = manifest["asset_catalog"]
    select_tool = SelectSlideTemplateSourceAssets(base_dir=_workspace(tmp_path))
    select_result = await select_tool.execute(
        ToolContext(tool_name=select_tool.name, arguments={}),
        SelectSlideTemplateSourceAssetsParams(
            template_dir=create_result.data["template_id"],
            decisions=[
                SourceAssetDecisionParams(
                    asset=catalog[0]["asset"],
                    decision="reuse",
                    asset_type="illustration",
                    quality="high",
                    rationale="The original artwork is the source visual anchor and will be used on the cover.",
                    target_name="cover-illustration.png",
                    intended_role="Cover visual anchor.",
                ),
                SourceAssetDecisionParams(
                    asset=catalog[1]["asset"],
                    decision="discard",
                    asset_type="other",
                    quality="poor",
                    rationale="The package-only image is not linked from the source slide and will not be used.",
                ),
            ],
        ),
    )
    assert select_result.ok, select_result.content
    await async_write_text(template_dir / "visual-spec.md", "# Visual Specification\n\nReuse source artwork.\n")
    await async_write_text(
        template_dir / "theme.css",
        "html, body { margin: 0; width: 1920px; height: 1080px; overflow: hidden; }\n.slide { width: 1920px; height: 1080px; }\n",
    )
    slide_without_asset = "<!doctype html><html><head><link rel=\"stylesheet\" href=\"../theme.css\"></head><body><main class=\"slide\"><h1 data-slot=\"title\">Business Review</h1><p data-slot=\"subtitle\">Q4</p><span data-slot=\"meta\">Strategy Team</span></main></body></html>"
    await async_write_text(template_dir / "slides/cover.html", slide_without_asset)

    from app.tools.slide_template_builder.common import validate_template_project

    with pytest.raises(Exception, match="Every selected source asset must be referenced"):
        await validate_template_project(template_dir)

    selected_manifest = await async_read_json(manifest_path)
    final_asset = next(
        item["final_file"]
        for item in selected_manifest["asset_catalog"]
        if item["decision"] == "reuse"
    )
    slide_with_asset = slide_without_asset.replace(
        "</main>",
        f'<img src="../{final_asset}" alt="Source illustration"></main>',
    )
    await async_write_text(template_dir / "slides/cover.html", slide_with_asset)
    validated = await validate_template_project(template_dir)
    assert validated.source_analysis_summary["referenced_source_asset_count"] == 1

    (template_dir / "images/source-assets/unused.png").write_bytes(b"unused")
    with pytest.raises(Exception, match="Unused files remain in images/source-assets/"):
        await validate_template_project(template_dir)


@pytest.mark.asyncio
async def test_finalize_project_generates_previews_beside_source(tmp_path):
    from playwright.async_api import async_playwright

    async with async_playwright() as playwright:
        if not Path(playwright.chromium.executable_path).exists():
            pytest.skip("Playwright Chromium is not installed in this environment")
        try:
            browser = await playwright.chromium.launch(headless=True)
            await browser.close()
        except Exception as exc:
            pytest.skip(f"Playwright Chromium cannot launch in this environment: {exc}")
    create_tool = CreateSlideTemplateProject(base_dir=_workspace(tmp_path))
    create_result = await create_tool.execute(ToolContext(tool_name=create_tool.name, arguments={}), _params())
    assert create_result.ok
    template_dir = Path(create_result.data["template_directory"])
    await async_write_text(
        template_dir / "visual-spec.md",
        "# Visual Specification\n\nNavy report system.\n",
    )
    await async_write_text(
        template_dir / "theme.css",
        ":root { --navy: #102a43; }\nhtml, body { margin: 0; width: 1920px; height: 1080px; overflow: hidden; }\n.slide { width: 1920px; height: 1080px; background: var(--navy); }\n",
    )
    await async_write_text(
        template_dir / "slides/cover.html",
        "<!doctype html><html><head><link rel=\"stylesheet\" href=\"../theme.css\"></head><body><main class=\"slide\"><h1 data-slot=\"title\" data-content-hint=\"Report title\">Business Review</h1><p data-slot=\"subtitle\" data-content-hint=\"Period\">Q4</p><span data-slot=\"meta\" data-content-hint=\"Presenter\">Strategy Team</span></main></body></html>",
    )

    finalize_tool = FinalizeSlideTemplate(base_dir=_workspace(tmp_path))
    result = await finalize_tool.execute(
        ToolContext(tool_name=finalize_tool.name, arguments={}),
        FinalizeSlideTemplateParams(template_dir=create_result.data["template_id"]),
    )

    assert result.ok, result.content
    work_dir = Path(create_result.data["working_directory"])
    assert not work_dir.exists()
    assert not (template_dir / WORK_DIR_NAME).exists()
    assert (template_dir / "previews/thumbnail.png").exists()
    assert (template_dir / "previews/collage.png").exists()
    assert (template_dir / "previews/slides/slide-001.png").exists()
    assert (template_dir / "template-assets.json").exists()
    with Image.open(template_dir / "previews/thumbnail.png") as image:
        assert image.size == (1280, 720)
    with Image.open(template_dir / "previews/slides/slide-001.png") as image:
        assert image.size == (1920, 1080)
    manifest = await async_read_json(template_dir / "template-assets.json")
    assert manifest["category_code"] == "PPT-CATE-business-report"
    assert manifest["tag_codes"] == _params().tag_codes
    assert manifest["local_files"]["thumbnail"] == "previews/thumbnail.png"
    assert manifest["backend_fields"]["thumbnail_file_key"] is None
    assert manifest["backend_fields"]["preview_image_file_keys"] == []
    assert manifest["backend_fields"]["status"] == 1
    assert manifest["backend_fields"]["category_code"] == "PPT-CATE-business-report"
    assert manifest["backend_fields"]["tag_codes"] == _params().tag_codes
    assert str(work_dir) in manifest["cleanup"]["removed_temporary_paths"]
    assert manifest["not_generated"]["template_file_key"]


def test_resolve_template_working_directory_uses_sibling_runtime_for_workspace(tmp_path):
    workspace = tmp_path / ".workspace"
    template_dir = workspace / "SLIDE-demo"
    template_id = "SLIDE-demo-ABCDEF"
    work_dir = resolve_template_working_directory(template_dir, template_id)

    assert work_dir.resolve() == (tmp_path / ".runtime" / "slide-template-maker" / template_id).resolve()


def test_resolve_template_working_directory_stays_outside_isolated_workspace(tmp_path):
    template_dir = tmp_path / "SLIDE-demo"
    template_id = "SLIDE-demo-ABCDEF"
    work_dir = resolve_template_working_directory(template_dir, template_id)

    assert work_dir.resolve() == (
        tmp_path.parent / f"{tmp_path.name}.runtime" / "slide-template-maker" / template_id
    ).resolve()
    with pytest.raises(ValueError):
        work_dir.resolve().relative_to(tmp_path.resolve())


@pytest.mark.asyncio
async def test_cleanup_template_intermediates_removes_runtime_and_legacy_work_dirs(tmp_path):
    template_dir = _workspace(tmp_path) / "SLIDE-navy-quarterly-review"
    template_id = "SLIDE-navy-quarterly-review-ABC123"
    template_dir.mkdir()
    (template_dir / "template.json").write_text(json.dumps({"template_id": template_id}), encoding="utf-8")
    work_dir = resolve_template_working_directory(template_dir, template_id)
    legacy_dir = template_dir / WORK_DIR_NAME
    final_images_dir = template_dir / "images"
    work_dir.mkdir(parents=True)
    legacy_dir.mkdir()
    final_images_dir.mkdir()
    (work_dir / "reference-render.png").write_bytes(b"temporary")
    (legacy_dir / "old-scratch.png").write_bytes(b"legacy")
    (final_images_dir / "hero.webp").write_bytes(b"final")

    removed = await cleanup_template_intermediates(template_dir)

    assert WORK_DIR_NAME in removed
    assert str(work_dir) in removed
    assert not work_dir.exists()
    assert not legacy_dir.exists()
    assert (final_images_dir / "hero.webp").exists()


def test_template_id_uses_backend_compatible_format():
    assert validate_template_id("SLIDE-navy-quarterly-review") == "SLIDE-navy-quarterly-review"
    assert validate_template_id("SLIDE-BUSINESS-MINIMAL") == "SLIDE-BUSINESS-MINIMAL"
    with pytest.raises(Exception):
        validate_template_id("SLIDE-中文模板")
    with pytest.raises(Exception):
        validate_template_id("PPT-navy-quarterly-review")


@pytest.mark.asyncio
async def test_evict_stale_template_working_directories_removes_old_dirs(tmp_path):
    root = tmp_path / ".runtime" / "slide-template-maker"
    old_dir = root / "old-id"
    fresh_dir = root / "fresh-id"
    old_dir.mkdir(parents=True)
    fresh_dir.mkdir()
    (old_dir / "scratch.png").write_bytes(b"old")
    (fresh_dir / "scratch.png").write_bytes(b"new")
    os.utime(old_dir, (0, 0))

    await evict_stale_template_working_directories(root, max_age_seconds=60, max_entries=8)

    assert not old_dir.exists()
    assert fresh_dir.exists()


def test_preview_helpers_keep_expected_dimensions():
    source = BytesIO()
    Image.new("RGB", (1920, 1080), "#123456").save(source, format="PNG")
    thumbnail = _resize_png(source.getvalue(), (1280, 720))
    collage = _create_collage([source.getvalue()])
    with Image.open(BytesIO(thumbnail)) as image:
        assert image.size == (1280, 720)
    with Image.open(BytesIO(collage)) as image:
        assert image.size == (704, 424)
