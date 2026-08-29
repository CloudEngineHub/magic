"""Read PPTX structure and embedded images without generating slide HTML."""

from __future__ import annotations

import hashlib
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.tools.slide_template_builder.canvas import normalize_canvas_size

SUPPORTED_IMAGE_EXTENSIONS = {
    "bmp",
    "emf",
    "gif",
    "jpeg",
    "jpg",
    "png",
    "svg",
    "tif",
    "tiff",
    "webp",
    "wmf",
}


@dataclass(frozen=True)
class ExtractedPptxAsset:
    relative_path: str
    sha256: str
    data: bytes


@dataclass(frozen=True)
class ExtractedPptxStructure:
    metadata: dict[str, Any]
    assets: list[ExtractedPptxAsset]


def _enum_name(value: Any) -> str | None:
    if value is None:
        return None
    name = getattr(value, "name", None)
    return str(name or value)


def _color_value(color: Any) -> dict[str, Any] | None:
    if color is None:
        return None
    result: dict[str, Any] = {}
    try:
        if color.rgb is not None:
            result["rgb"] = str(color.rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    try:
        if color.theme_color is not None:
            result["theme"] = _enum_name(color.theme_color)
    except (AttributeError, TypeError, ValueError):
        pass
    try:
        if color.brightness is not None:
            result["brightness"] = float(color.brightness)
    except (AttributeError, TypeError, ValueError):
        pass
    return result or None


def _fill_value(fill: Any) -> dict[str, Any] | None:
    if fill is None:
        return None
    fill_type = _enum_name(getattr(fill, "type", None))
    result: dict[str, Any] = {"type": fill_type} if fill_type else {}
    try:
        color = _color_value(fill.fore_color)
        if color:
            result["color"] = color
    except (AttributeError, TypeError, ValueError):
        pass
    return result or None


def _line_value(line: Any) -> dict[str, Any] | None:
    if line is None:
        return None
    result: dict[str, Any] = {}
    width = getattr(line, "width", None)
    if width is not None:
        result["width_pt"] = round(float(width.pt), 2)
    try:
        color = _color_value(line.color)
        if color:
            result["color"] = color
    except (AttributeError, TypeError, ValueError):
        pass
    dash_style = _enum_name(getattr(line, "dash_style", None))
    if dash_style:
        result["dash_style"] = dash_style
    return result or None


def _font_value(font: Any) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key in ("name", "bold", "italic", "underline"):
        value = getattr(font, key, None)
        if value is not None:
            result[key] = value
    size = getattr(font, "size", None)
    if size is not None:
        result["size_pt"] = round(float(size.pt), 2)
    try:
        color = _color_value(font.color)
        if color:
            result["color"] = color
    except (AttributeError, TypeError, ValueError):
        pass
    return result


def _text_value(shape: Any) -> dict[str, Any] | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    paragraphs = []
    for paragraph in shape.text_frame.paragraphs:
        runs = []
        for run in paragraph.runs:
            font = _font_value(run.font)
            if run.text or font:
                runs.append({"text": run.text, "font": font})
        paragraphs.append({
            "text": paragraph.text,
            "alignment": _enum_name(paragraph.alignment),
            "level": paragraph.level,
            "runs": runs,
        })
    margins = {}
    for key in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        value = getattr(shape.text_frame, key, None)
        if value is not None:
            margins[key] = int(value)
    return {
        "text": shape.text,
        "vertical_anchor": _enum_name(shape.text_frame.vertical_anchor),
        "word_wrap": shape.text_frame.word_wrap,
        "margins_emu": margins,
        "paragraphs": paragraphs,
    }


def _bbox(
    shape: Any,
    slide_width: int,
    slide_height: int,
    normalized_width: int,
    normalized_height: int,
) -> dict[str, Any]:
    left = int(getattr(shape, "left", 0) or 0)
    top = int(getattr(shape, "top", 0) or 0)
    width = int(getattr(shape, "width", 0) or 0)
    height = int(getattr(shape, "height", 0) or 0)
    return {
        "emu": {"left": left, "top": top, "width": width, "height": height},
        "canvas": {
            "left": round(left / slide_width * normalized_width, 2),
            "top": round(top / slide_height * normalized_height, 2),
            "width": round(width / slide_width * normalized_width, 2),
            "height": round(height / slide_height * normalized_height, 2),
        },
    }


def _table_value(shape: Any) -> dict[str, Any] | None:
    if not getattr(shape, "has_table", False):
        return None
    table = shape.table
    return {
        "rows": len(table.rows),
        "columns": len(table.columns),
        "cells": [[cell.text for cell in row.cells] for row in table.rows],
    }


def _related_picture_blob(shape: Any) -> tuple[bytes | None, str, str | None]:
    """Read picture bytes when python-pptx cannot expose an ImagePart."""
    try:
        pic = getattr(shape, "_pic", None)
        r_id = pic.blip_rId if pic is not None else None
        if not r_id:
            return None, "bin", None
        part = shape.part.related_part(r_id)
        blob = getattr(part, "blob", None)
        if blob is None:
            return None, "bin", None
        partname = str(getattr(part, "partname", "") or "")
        extension = Path(partname).suffix.lower().lstrip(".") or "bin"
        content_type = getattr(part, "content_type", None)
        return blob, extension, str(content_type) if content_type else None
    except (AttributeError, KeyError, TypeError, ValueError):
        return None, "bin", None


def _picture_crop(shape: Any) -> dict[str, float] | None:
    try:
        return {
            "left": float(shape.crop_left),
            "right": float(shape.crop_right),
            "top": float(shape.crop_top),
            "bottom": float(shape.crop_bottom),
        }
    except (AttributeError, TypeError, ValueError):
        return None


def _picture_value(
    shape: Any,
    asset_paths: dict[str, str],
    asset_blobs: dict[str, bytes],
) -> dict[str, Any] | None:
    blob = None
    extension = "bin"
    content_type = None
    source_size = None
    try:
        image = shape.image
        blob = image.blob
        extension = str(image.ext or "bin").lower()
        content_type = image.content_type
        source_size = {"width": image.size[0], "height": image.size[1]}
    except (AttributeError, KeyError, TypeError, ValueError):
        blob, extension, content_type = _related_picture_blob(shape)
    if not blob:
        return None
    digest = hashlib.sha256(blob).hexdigest()
    relative_path = asset_paths.setdefault(digest, f"source-assets/{digest[:16]}.{extension}")
    asset_blobs.setdefault(digest, blob)
    result: dict[str, Any] = {
        "asset": relative_path,
        "sha256": digest,
    }
    if content_type:
        result["content_type"] = content_type
    if source_size:
        result["source_size"] = source_size
    crop = _picture_crop(shape)
    if crop:
        result["crop"] = crop
    return result


def _shape_value(
    shape: Any,
    slide_width: int,
    slide_height: int,
    normalized_width: int,
    normalized_height: int,
    asset_paths: dict[str, str],
    asset_blobs: dict[str, bytes],
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "name": shape.name,
        "type": _enum_name(shape.shape_type),
        "bbox": _bbox(shape, slide_width, slide_height, normalized_width, normalized_height),
        "rotation": float(getattr(shape, "rotation", 0) or 0),
    }
    if getattr(shape, "is_placeholder", False):
        try:
            result["placeholder_type"] = _enum_name(shape.placeholder_format.type)
        except (AttributeError, KeyError, ValueError):
            pass
    text = _text_value(shape)
    if text:
        result["text"] = text
    fill = _fill_value(getattr(shape, "fill", None))
    if fill:
        result["fill"] = fill
    line = _line_value(getattr(shape, "line", None))
    if line:
        result["line"] = line
    table = _table_value(shape)
    if table:
        result["table"] = table
    if getattr(shape, "has_chart", False):
        result["chart"] = {"type": _enum_name(shape.chart.chart_type)}
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image = _picture_value(shape, asset_paths, asset_blobs)
        if image:
            result["image"] = image
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        result["children"] = [
            _shape_value(
                child,
                slide_width,
                slide_height,
                normalized_width,
                normalized_height,
                asset_paths,
                asset_blobs,
            )
            for child in shape.shapes
        ]
    return result


def _collect_package_images(
    source_path: Path,
    asset_paths: dict[str, str],
    asset_blobs: dict[str, bytes],
) -> None:
    with zipfile.ZipFile(source_path) as archive:
        for member in archive.namelist():
            if not member.startswith("ppt/media/") or member.endswith("/"):
                continue
            extension = Path(member).suffix.lower().lstrip(".")
            if extension not in SUPPORTED_IMAGE_EXTENSIONS:
                continue
            blob = archive.read(member)
            digest = hashlib.sha256(blob).hexdigest()
            asset_paths.setdefault(digest, f"source-assets/{digest[:16]}.{extension}")
            asset_blobs.setdefault(digest, blob)


def extract_pptx_structure(source_path: Path) -> ExtractedPptxStructure:
    presentation = Presentation(str(source_path))
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    normalized_canvas = normalize_canvas_size(slide_width, slide_height)
    asset_paths: dict[str, str] = {}
    asset_blobs: dict[str, bytes] = {}
    slides = []
    for slide_index, slide in enumerate(presentation.slides, 1):
        shapes = [
            _shape_value(
                shape,
                slide_width,
                slide_height,
                normalized_canvas.width,
                normalized_canvas.height,
                asset_paths,
                asset_blobs,
            )
            for shape in slide.shapes
        ]
        background = _fill_value(slide.background.fill)
        slides.append({
            "index": slide_index,
            "background": background,
            "shape_count": len(shapes),
            "shapes": shapes,
        })
    _collect_package_images(source_path, asset_paths, asset_blobs)
    assets = [
        ExtractedPptxAsset(relative_path=asset_paths[digest], sha256=digest, data=blob)
        for digest, blob in asset_blobs.items()
    ]
    return ExtractedPptxStructure(
        metadata={
            "schema_version": "1.0",
            "source_file": source_path.name,
            "slide_size_emu": {"width": slide_width, "height": slide_height},
            "normalized_canvas": normalized_canvas.model_dump(),
            "slide_count": len(slides),
            "source_asset_count": len(assets),
            "slides": slides,
        },
        assets=assets,
    )
